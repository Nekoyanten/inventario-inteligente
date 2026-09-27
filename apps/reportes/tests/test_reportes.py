import io

import pytest
from openpyxl import Workbook, load_workbook

from apps.catalogo.models import Producto
from apps.inventario.models import TipoMovimiento
from apps.inventario.services import registrar_movimiento
from apps.reportes.reportes import REPORTES
from apps.usuarios.models import Rol, Usuario
from apps.ventas.services import registrar_venta


@pytest.fixture
def con_datos(producto, admin):
    registrar_movimiento(producto=producto, tipo=TipoMovimiento.ENTRADA_INICIAL, cantidad=20, usuario=admin)
    registrar_venta(negocio=producto.negocio, vendedor=admin, lineas=[{"producto": producto, "cantidad": 3}])
    return producto


@pytest.mark.parametrize("clave", list(REPORTES))
def test_todos_los_reportes_se_ven_y_exportan(client, admin, con_datos, clave):
    client.force_login(admin)
    assert client.get(f"/reportes/{clave}/").status_code == 200
    for formato, firma in (("csv", None), ("xlsx", b"PK"), ("pdf", b"%PDF")):
        resp = client.get(f"/reportes/{clave}.{formato}")
        assert resp.status_code == 200, (clave, formato)
        if firma:
            assert resp.content[:len(firma)] == firma


def test_utilidad_calcula_margen(client, admin, con_datos):
    client.force_login(admin)
    tabla = client.get("/reportes/utilidad/").context["tabla"]
    fila = tabla.filas[0]
    assert fila[3] == 36000 and fila[4] == 25500 and fila[5] == 10500


def test_encargado_no_ve_reportes_financieros(client, negocio, con_datos):
    enc = Usuario.objects.create_user("e", password="x", negocio=negocio, rol=Rol.INVENTARIO)
    client.force_login(enc)
    assert client.get("/reportes/utilidad/").status_code == 403
    assert "Utilidad estimada" not in client.get("/reportes/").content.decode()
    assert client.get("/reportes/inventario-actual/").status_code == 200


def _xlsx(filas):
    wb = Workbook()
    for f in filas:
        wb.active.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "productos.xlsx"
    return buf


def test_importar_productos(client, admin, negocio):
    client.force_login(admin)
    assert load_workbook(io.BytesIO(client.get("/reportes/importar/plantilla.xlsx").content)).active["A1"].value == "sku"
    archivo = _xlsx([["sku", "nombre", "categoria", "precio_venta", "stock_inicial", "proveedor"],
                     ["jab-1", "Jabón", "Aseo", 3500, 12, "Proveedor Nuevo"],
                     ["LEC-1", "Leche", "Lácteos", "3.800", 0, ""]])
    resp = client.post("/reportes/importar/", {"archivo": archivo})
    assert resp.status_code == 302
    jabon = Producto.objects.get(sku="JAB-1")
    assert jabon.stock_actual == 12 and jabon.categoria.nombre == "Aseo" and jabon.proveedor_principal.nombre == "Proveedor Nuevo"


def test_importacion_con_errores_no_importa_nada(client, admin):
    client.force_login(admin)
    archivo = _xlsx([["sku", "nombre", "precio_venta"], ["A", "Bien", 100], ["", "Sin sku", "abc"]])
    html = client.post("/reportes/importar/", {"archivo": archivo}).content.decode()
    assert "Fila 3" in html and not Producto.objects.filter(sku="A").exists()
