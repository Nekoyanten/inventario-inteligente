"""Importar productos desde Excel o CSV (para negocios que ya tienen su lista en una hoja de cálculo)."""

import csv
import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import redirect, render

from apps.catalogo.models import Categoria, Marca, Producto, UnidadMedida
from apps.core.auditoria import auditar
from apps.core.limites import LimiteDelPlan, verificar_productos
from apps.core.negocio import negocio_requerido
from apps.inventario.models import TipoMovimiento
from apps.inventario.services import ErrorInventario, registrar_movimiento
from apps.proveedores.models import Proveedor
from apps.usuarios.permisos import requiere_permiso

COLUMNAS = [
    ("sku", "Código o SKU (obligatorio)"), ("nombre", "Nombre (obligatorio)"), ("categoria", "Categoría"),
    ("marca", "Marca"), ("unidad", "Unidad (und, kg, L…)"), ("precio_compra", "Costo"),
    ("precio_venta", "Precio de venta"), ("stock_minimo", "Stock mínimo"), ("stock_inicial", "Stock inicial"),
    ("codigo_barras", "Código de barras"), ("proveedor", "Proveedor"), ("vencimiento", "Vencimiento (AAAA-MM-DD)"),
    ("vida_util_dias", "Vida útil en días (perecederos)"),
]
EJEMPLO = ["ARR-1K", "Arroz 1 kg", "Abarrotes", "Diana", "und", 3200, 4500, 10, 50, "7702001001", "Distribuidora A",
           "2027-03-31", ""]
MAX_FILAS = 5000


def _entero_positivo(valor, errores):
    if valor in (None, ""):
        return None
    try:
        n = int(float(str(valor).replace(",", ".")))
    except ValueError:
        errores.append(f"vida útil inválida ({valor})")
        return None
    if n <= 0:
        errores.append("la vida útil debe ser mayor que cero")
        return None
    return n


def plantilla(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Productos"
    ws.append([c for c, _ in COLUMNAS])
    ws.append(EJEMPLO)
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F7A4D")
    ayuda = wb.create_sheet("Instrucciones")
    ayuda.append(["Columna", "Descripción"])
    for c, d in COLUMNAS:
        ayuda.append([c, d])
    ayuda.append([])
    ayuda.append(["Borra la fila de ejemplo antes de importar. Si el SKU ya existe se actualizan sus datos "
                  "(el stock inicial solo se aplica a productos nuevos)."])
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = 'attachment; filename="plantilla-productos.xlsx"'
    return resp


def leer_archivo(archivo) -> list[dict]:
    from django.conf import settings

    if archivo.size > settings.TAMANO_MAX_ARCHIVO_MB * 1024 * 1024:
        raise ValueError(f"El archivo supera {settings.TAMANO_MAX_ARCHIVO_MB} MB.")
    nombre = archivo.name.lower()
    if nombre.endswith(".csv"):
        texto = archivo.read().decode("utf-8-sig")
        separador = ";" if texto.count(";") > texto.count(",") else ","
        filas = list(csv.reader(io.StringIO(texto), delimiter=separador))
    elif nombre.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(archivo, read_only=True, data_only=True)
        filas = [list(f) for f in wb.worksheets[0].iter_rows(values_only=True)]
    else:
        raise ValueError("Sube un archivo .xlsx o .csv.")
    if not filas:
        raise ValueError("El archivo está vacío.")
    encabezados = [str(h or "").strip().lower() for h in filas[0]]
    if "sku" not in encabezados or "nombre" not in encabezados:
        raise ValueError("El archivo debe tener las columnas «sku» y «nombre». Descarga la plantilla.")
    datos = []
    for fila in filas[1:MAX_FILAS + 1]:
        if not any(v not in (None, "") for v in fila):
            continue
        datos.append({encabezados[i]: fila[i] for i in range(min(len(encabezados), len(fila)))})
    return datos


def _decimal(valor, campo, errores, defecto="0"):
    if valor in (None, ""):
        return Decimal(defecto)
    try:
        d = Decimal(str(valor).replace("$", "").replace(" ", "").replace(",", "."))
        if d < 0:
            raise InvalidOperation
        return d
    except InvalidOperation:
        errores.append(f"«{campo}» no es un número válido ({valor})")
        return Decimal(defecto)


def _fecha(valor, errores):
    if valor in (None, ""):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    try:
        return date.fromisoformat(str(valor).strip()[:10])
    except ValueError:
        errores.append(f"vencimiento inválido ({valor}); usa AAAA-MM-DD")
        return None


def validar(negocio, filas: list[dict]) -> tuple[list[dict], list[str]]:
    """Devuelve (filas_limpias, errores). Si hay errores no se importa nada."""
    limpias, errores, vistos = [], [], set()
    unidades = {u.abreviatura.lower(): u for u in UnidadMedida.objects.all()} | {
        u.nombre.lower(): u for u in UnidadMedida.objects.all()}
    for n, fila in enumerate(filas, start=2):
        e = []
        sku = str(fila.get("sku") or "").strip().upper()
        nombre = str(fila.get("nombre") or "").strip()
        if not sku:
            e.append("falta el SKU")
        if not nombre:
            e.append("falta el nombre")
        if sku in vistos:
            e.append(f"el SKU {sku} está repetido en el archivo")
        vistos.add(sku)
        unidad_txt = str(fila.get("unidad") or "").strip().lower()
        unidad = unidades.get(unidad_txt) if unidad_txt else None
        if unidad_txt and unidad is None:
            e.append(f"unidad desconocida ({unidad_txt})")
        limpia = {
            "sku": sku[:40], "nombre": nombre[:150], "categoria": str(fila.get("categoria") or "").strip()[:80],
            "marca": str(fila.get("marca") or "").strip()[:80], "unidad": unidad,
            "precio_compra": _decimal(fila.get("precio_compra"), "precio_compra", e),
            "precio_venta": _decimal(fila.get("precio_venta"), "precio_venta", e),
            "stock_minimo": _decimal(fila.get("stock_minimo"), "stock_minimo", e),
            "stock_inicial": _decimal(fila.get("stock_inicial"), "stock_inicial", e),
            "codigo_barras": str(fila.get("codigo_barras") or "").strip()[:40],
            "proveedor": str(fila.get("proveedor") or "").strip()[:150],
            "vencimiento": _fecha(fila.get("vencimiento"), e),
            "vida_util_dias": _entero_positivo(fila.get("vida_util_dias"), e),
        }
        if e:
            errores.append(f"Fila {n}: " + "; ".join(e))
        limpias.append(limpia)
    return limpias, errores


@transaction.atomic
def importar_filas(negocio, usuario, filas: list[dict]) -> dict:
    creados = actualizados = 0
    cache_cat, cache_marca, cache_prov = {}, {}, {}

    def obtener(cache, modelo, nombre):
        if not nombre:
            return None
        if nombre.lower() not in cache:
            cache[nombre.lower()] = modelo.objects.filter(negocio=negocio, nombre__iexact=nombre).first() or \
                modelo.objects.create(negocio=negocio, nombre=nombre)
        return cache[nombre.lower()]

    for f in filas:
        datos = {
            "nombre": f["nombre"], "categoria": obtener(cache_cat, Categoria, f["categoria"]),
            "marca": obtener(cache_marca, Marca, f["marca"]), "precio_compra": f["precio_compra"],
            "precio_venta": f["precio_venta"], "stock_minimo": f["stock_minimo"], "codigo_barras": f["codigo_barras"],
        }
        if f.get("vida_util_dias"):
            datos["vida_util_dias"] = f["vida_util_dias"]
        if f["unidad"]:
            datos["unidad"] = f["unidad"]
        proveedor = obtener(cache_prov, Proveedor, f["proveedor"])
        if proveedor:
            datos["proveedor_principal"] = proveedor
        producto, creado = Producto.objects.update_or_create(negocio=negocio, sku=f["sku"], defaults=datos)
        if creado:
            creados += 1
            if f["stock_inicial"] > 0:
                registrar_movimiento(producto=producto, tipo=TipoMovimiento.ENTRADA_INICIAL, cantidad=f["stock_inicial"],
                                     usuario=usuario, motivo="Importación desde archivo",
                                     fecha_vencimiento=f["vencimiento"], evaluar_alertas=False)
        else:
            actualizados += 1
    auditar(negocio, usuario, "importar_productos", negocio, creados=creados, actualizados=actualizados)
    return {"creados": creados, "actualizados": actualizados}


@negocio_requerido
@requiere_permiso("gestionar_productos")
def importar(request):
    errores = []
    if request.method == "POST" and request.FILES.get("archivo"):
        try:
            filas, errores = validar(request.negocio, leer_archivo(request.FILES["archivo"]))
            if not filas:
                errores = ["El archivo no tiene productos."]
            if not errores:
                existentes = set(Producto.objects.filter(negocio=request.negocio).values_list("sku", flat=True))
                verificar_productos(request.negocio, sum(1 for f in filas if f["sku"] not in existentes))
                r = importar_filas(request.negocio, request.user, filas)
                messages.success(request, f"Importación lista: {r['creados']} productos nuevos y "
                                          f"{r['actualizados']} actualizados.")
                return redirect("catalogo:lista")
        except (ValueError, ErrorInventario, LimiteDelPlan) as e:
            errores = [str(e)]
    return render(request, "reportes/importar.html", {"errores": errores[:50], "mas_errores": max(0, len(errores) - 50),
                                                      "columnas": COLUMNAS})
