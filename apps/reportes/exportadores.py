"""Exportación de reportes a CSV, Excel y PDF."""

import csv
import io
from datetime import date, datetime
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone


def _plano(valor):
    if isinstance(valor, datetime):
        return timezone.localtime(valor).strftime("%Y-%m-%d %H:%M") if timezone.is_aware(valor) else valor.strftime(
            "%Y-%m-%d %H:%M")
    if isinstance(valor, date):
        return valor.isoformat()
    if valor is None:
        return ""
    return valor


def a_csv(nombre: str, encabezados: list[str], filas) -> HttpResponse:
    resp = HttpResponse(content_type="text/csv; charset=utf-8")
    resp["Content-Disposition"] = f'attachment; filename="{nombre}.csv"'
    resp.write("﻿")  # BOM para que Excel abra bien las tildes
    w = csv.writer(resp)
    w.writerow(encabezados)
    w.writerows([[_plano(v) for v in f] for f in filas])
    return resp


def a_excel(nombre: str, encabezados: list[str], filas, tipos=None, totales=None, titulo=None) -> HttpResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = (titulo or nombre)[:31]
    ws.append(encabezados)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F7A4D")
    formatos = {"moneda": '"$"#,##0', "numero": "#,##0.##", "porcentaje": '0.0"%"', "fecha": "yyyy-mm-dd"}
    tipos = tipos or []
    for fila in list(filas) + ([totales] if totales else []):
        valores = []
        for v in fila:
            if isinstance(v, Decimal):
                v = float(v)
            elif isinstance(v, datetime) and timezone.is_aware(v):
                v = timezone.localtime(v).replace(tzinfo=None)
            valores.append(v)
        ws.append(valores)
        for i, t in enumerate(tipos):
            if t in formatos:
                ws.cell(row=ws.max_row, column=i + 1).number_format = formatos[t]
    if totales:
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
    for i, enc in enumerate(encabezados, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(45, len(enc) + 6))
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{nombre}.xlsx"'
    return resp


def a_pdf(nombre: str, titulo: str, subtitulo: str, encabezados: list[str], filas, tipos=None, totales=None) -> HttpResponse:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from apps.core.formato import numero

    def fmt(v, t):
        if v in (None, ""):
            return ""
        if t == "moneda":
            return "$" + f"{float(v):,.0f}".replace(",", ".")
        if t == "porcentaje":
            return f"{float(v):.1f} %"
        if t == "numero" and isinstance(v, (int, float, Decimal)):
            return numero(v)
        v = _plano(v)
        return str(v)[:60]

    tipos = tipos or ["texto"] * len(encabezados)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(letter), leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                            topMargin=1.5 * cm, bottomMargin=1.5 * cm, title=titulo)
    est = getSampleStyleSheet()
    datos = [encabezados] + [[fmt(v, tipos[i]) for i, v in enumerate(f)] for f in filas]
    if totales:
        datos.append([fmt(v, tipos[i]) for i, v in enumerate(totales)])
    tabla = Table(datos, repeatRows=1)
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f7a4d")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#dddddd")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f6f5")]),
    ]
    for i, t in enumerate(tipos):
        if t in ("moneda", "numero", "porcentaje"):
            estilo.append(("ALIGN", (i, 0), (i, -1), "RIGHT"))
    if totales:
        estilo.append(("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"))
    tabla.setStyle(TableStyle(estilo))
    doc.build([Paragraph(titulo, est["Title"]), Paragraph(subtitulo, est["Normal"]), Spacer(1, 10), tabla])
    resp = HttpResponse(buf.getvalue(), content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="{nombre}.pdf"'
    return resp
